"""Write a Document as a Word file (.docx): title page, a contents list that links to each part, headings that
show in Word's navigation pane, and page numbers."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document as NewDocx
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from .build import Document, Section
from .md import Block, Run

INK = RGBColor(0x1C, 0x2A, 0x3A)
INK2 = RGBColor(0x44, 0x54, 0x6A)
MUTED = RGBColor(0x75, 0x83, 0x94)
TEAL = RGBColor(0x1F, 0x7A, 0x74)
SANS = "Calibri"
MAX_TOC_FILES = 400
MONO = "Consolas"
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")  # characters a .docx can't hold


def _para(d, style: str | None = None) -> Paragraph:
    """A new paragraph at the end of the document. Faster than add_paragraph, which searches the whole body."""
    p = OxmlElement("w:p")
    cache = d.__dict__.setdefault("_reader_cache", {})
    if "sect" not in cache:
        cache["sect"] = d.element.body.sectPr  # the body's last element; new paragraphs go before it
        cache["body"] = d._body
    if cache["sect"] is not None:
        cache["sect"].addprevious(p)
    else:
        d.element.body.append(p)
    para = Paragraph(p, cache["body"])
    if style:
        if style not in cache:
            cache[style] = d.styles[style].style_id
        para._p.style = cache[style]
    return para


def _t(text: str) -> str:
    return _BAD_XML.sub("", text)


def _font(run, name: str) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def _shade(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _bottom_border(paragraph, color: str = "D3DCE4") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "4"), ("w:color", color)):
        bottom.set(qn(k), v)
    borders.append(bottom)
    ppr.append(borders)


def _bookmark(paragraph, name: str, bid: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def _hyperlink(paragraph, text: str, *, url: str | None = None, anchor: str | None = None, mono: bool = False,
               color: RGBColor = TEAL, size: float | None = None, bold: bool = False):
    link = OxmlElement("w:hyperlink")
    if url:
        link.set(qn("r:id"), paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/"
                                                           "relationships/hyperlink", is_external=True))
    else:
        link.set(qn("w:anchor"), anchor)
    run = paragraph.add_run(_t(text))
    run.font.color.rgb = color
    if url:
        run.font.underline = True
    if mono:
        _font(run, MONO)
    if size:
        run.font.size = Pt(size)
    if bold:
        run.bold = True
    link.append(run._r)
    paragraph._p.append(link)
    return run


def _page_number(paragraph) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _runs(paragraph, runs: list[Run], size: float = 10.5) -> None:
    for r in runs:
        if r.link:
            run = _hyperlink(paragraph, r.text, url=r.link, mono=r.code, size=size * (0.9 if r.code else 1))
        else:
            run = paragraph.add_run(_t(r.text))
        if r.bold:
            run.bold = True
        if r.italic:
            run.italic = True
        if r.link:
            continue
        if r.code:
            _font(run, MONO)
            run.font.size = Pt(size * 0.9)
            run.font.color.rgb = INK2


def _blocks(d, blocks: list[Block]) -> None:
    for b in blocks:
        if b.kind == "p":
            _runs(_para(d), b.runs)
        elif b.kind == "h":
            p = _para(d)
            p.paragraph_format.keep_with_next = True
            _runs(p, b.runs)
            for run in p.runs:
                run.bold = True
        elif b.kind == "item":
            p = _para(d)
            fmt = p.paragraph_format
            fmt.left_indent = Mm(6 + 6 * b.level)
            fmt.first_line_indent = Mm(-4.5)
            fmt.space_after = Pt(2)
            marker = b.marker if b.marker != "•" else ("•" if b.level == 0 else "◦")
            p.add_run(marker + "\t")
            fmt.tab_stops.add_tab_stop(Mm(6 + 6 * b.level))
            _runs(p, b.runs)
        elif b.kind == "code":
            p = _para(d)
            p.paragraph_format.left_indent = Mm(2)
            p.paragraph_format.right_indent = Mm(2)
            _shade(p, "F2F5F8")
            for k, line in enumerate(_t(b.text).expandtabs(4).split("\n")):
                run = p.add_run(line)
                _font(run, MONO)
                run.font.size = Pt(8.5)
                run.font.color.rgb = INK2
                if k:
                    run._r.addprevious(_break())


def _break():
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:br"))
    return r


def _styles(d) -> None:
    st = d.styles
    normal = st["Normal"]
    normal.font.name = SANS
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, space_before in (("Heading 1", 18, 18), ("Heading 2", 13.5, 14), ("Heading 3", 11, 10)):
        h = st[name]
        h.font.name = SANS
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = INK
        h.paragraph_format.space_before = Pt(space_before)
        h.paragraph_format.space_after = Pt(1)
        h.paragraph_format.keep_with_next = True
        rpr = h.element.get_or_add_rPr()
        for el in rpr.findall(qn("w:rFonts")):
            rpr.remove(el)
    title = st["Title"]
    title.font.name = SANS
    title.font.size = Pt(30)
    title.font.color.rgb = INK
    title.font.bold = True


def _section(d, s: Section, bid: int) -> None:
    h = _para(d, f"Heading {s.level}")
    if s.level == 1:
        h.paragraph_format.page_break_before = False
        _bottom_border(h)
    run = h.add_run(_t(s.title))
    if s.mono:
        _font(run, MONO)
    _bookmark(h, s.anchor, bid)
    meta = _para(d)
    meta.paragraph_format.keep_with_next = True
    meta.paragraph_format.space_after = Pt(6)
    m = meta.add_run(_t(s.meta))
    m.font.size = Pt(8.5)
    m.font.color.rgb = MUTED
    if s.ai:
        a = meta.add_run(" · written by AI")
        a.font.size = Pt(8.5)
        a.font.color.rgb = TEAL
    _blocks(d, s.blocks)


def write_docx(doc: Document, path: str | Path, progress=None) -> None:
    d = NewDocx()
    _styles(d)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Mm(22))
    sec.top_margin = sec.bottom_margin = Mm(20)
    props = d.core_properties
    props.title = f"{doc.title}: the code, explained"
    props.author = "The Reader"
    props.subject = "A layered explanation of a codebase"

    # title page
    for _ in range(5):
        _para(d)
    t = _para(d, "Title")
    t.add_run(_t(doc.title))
    sub = _para(d)
    r = sub.add_run("The code, explained layer by layer")
    r.font.size = Pt(13)
    r.font.color.rgb = INK2
    sub.paragraph_format.space_after = Pt(24)
    table = d.add_table(rows=0, cols=2)
    for k, v in doc.facts:
        cells = table.add_row().cells
        kr = cells[0].paragraphs[0].add_run(k)
        kr.font.size = Pt(9)
        kr.font.color.rgb = MUTED
        vr = cells[1].paragraphs[0].add_run(_t(v))
        vr.font.size = Pt(9.5)
    table.autofit = False
    for col, width in zip(table.columns, (Mm(32), Mm(134))):
        col.width = width
        for cell in col.cells:
            cell.width = width
    for n in doc.notes:
        p = _para(d)
        p.paragraph_format.space_before = Pt(10)
        nr = p.add_run(_t(n))
        nr.italic = True
        nr.font.size = Pt(9)
        nr.font.color.rgb = INK2
    made = _para(d)
    made.paragraph_format.space_before = Pt(28)
    mr = made.add_run("Made with The Reader")
    mr.font.size = Pt(8.5)
    mr.font.color.rgb = MUTED

    # contents: the overview and folders, each a link to its part
    body = d.add_section(WD_SECTION.NEW_PAGE)
    footer = body.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fr = fp.add_run(_t(doc.title) + "    ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED
    _page_number(fp)
    for run in fp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
    sec.footer.is_linked_to_previous = False  # no footer on the title page
    c = _para(d)
    cr = c.add_run("Contents")
    cr.bold = True
    cr.font.size = Pt(18)
    c.paragraph_format.space_after = Pt(12)
    with_files = sum(1 for s in doc.sections if s.level == 2) <= MAX_TOC_FILES
    for s in doc.sections:
        if s.level == 1 or (s.level == 2 and with_files):
            p = _para(d)
            p.paragraph_format.space_after = Pt(1)
            if s.level == 1:
                p.paragraph_format.space_before = Pt(6)
            else:
                p.paragraph_format.left_indent = Mm(5)
            _hyperlink(p, s.title, anchor=s.anchor, mono=s.mono, color=INK if s.level == 1 else INK2,
                       size=10 if s.level == 1 else 8.5, bold=s.level == 1)
    _para(d).add_run().add_break(WD_BREAK.PAGE)

    for i, s in enumerate(doc.sections):
        _section(d, s, i + 1)
    d.save(str(path))
